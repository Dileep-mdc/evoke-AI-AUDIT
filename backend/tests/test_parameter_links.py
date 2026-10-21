"""The per-company workbook of which links each parameter checked, and why (app/parameter_links.py).

Pinned here: the file is named after the company in the URL; handlers record exactly the links
they score, with a reason, from their own selection (service pages only for ON-05, no utility
pages for ON-19, counted vs not-counted search results off-page); the links never enter the
evidence the model grades; and a re-check rewrites the workbook without losing the parameters
it did not re-run.
"""
from __future__ import annotations

import asyncio
import json
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

import app.crawler.google_search as gs
from app import parameter_links
from app.parameters import offpage, offpage_tools
from app.parameters.common import CHECKED_LINKS, link, merge_links
from app.parameters.engine import load_registry, run_one
from app.parameters.onpage import on_05, on_19
from app.parameters.technical import tech_01

SPECS = {p["parameter_id"]: p for p in load_registry()}


def _run(coro):
    return asyncio.run(coro)


def _page(url, *, page_type="other", words=300, text="Some page text about the services we deliver.", title="T"):
    return SimpleNamespace(
        url=url, result=SimpleNamespace(final_url=url, ok=True, text="", status_code=200, error=None, hops=0, elapsed_ms=50, content=b""),
        soup=None, text=text, word_count=words, title=title, meta_description="", canonical="", page_type=page_type,
        headings=[], links=[], images=[], schema_blocks=[], hreflang=[], dates={}, blocked_reason="")


def _ctx(pages, **extra):
    return SimpleNamespace(origin="https://www.acme.com", domain="www.acme.com", pages=pages,
                           homepage=pages[0] if pages else None, company_name="Acme Corp", brand_terms=["acme"],
                           crawl_errors=[], **extra)


@pytest.mark.parametrize("url, name", [
    ("https://www.evoketechnologies.com/", "evoketechnologies"),
    ("www.evoketechnologies.com", "evoketechnologies"),
    ("https://en.nvent.com/products", "nvent"),
    ("https://shop.example.co.uk/", "example"),
    ("https://aaos.org", "aaos"),
    ("http://localhost:8000/", "localhost"),
])
def test_the_workbook_is_named_after_the_company_in_the_url(url, name):
    assert parameter_links.company_from_url(url) == name


def test_a_link_read_for_two_reasons_is_listed_once_with_both():
    merged = merge_links([link("https://a/", "x"), link("https://a/", "y"), link("https://a/", "x"), link("", "z"), None])
    assert merged == [{"url": "https://a/", "why": "x Also: y"}]


def test_tech01_records_the_robots_file_it_read():
    robots = SimpleNamespace(url="https://www.acme.com/robots.txt", result=SimpleNamespace(ok=False), raw="")
    row = _run(tech_01(SPECS["TECH-01"], _ctx([_page("https://www.acme.com/")], robots=robots)))
    assert [x["url"] for x in row[CHECKED_LINKS]] == ["https://www.acme.com/robots.txt"]
    assert "AI crawler" in row[CHECKED_LINKS][0]["why"]


def test_on05_lists_the_service_pages_it_scored_and_nothing_else():
    pages = [_page("https://www.acme.com/", page_type="home"), _page("https://www.acme.com/services/cloud", page_type="service"),
             _page("https://www.acme.com/blog/news"), _page("https://www.acme.com/contact", page_type="utility")]
    row = _run(on_05(SPECS["ON-05"], _ctx(pages)))
    assert [x["url"] for x in row[CHECKED_LINKS]] == ["https://www.acme.com/services/cloud"]
    assert row[CHECKED_LINKS][0]["why"].startswith("Service or solution page")
    assert row["evidence"]["service_pages_total"] == len(row[CHECKED_LINKS])


def test_on05_says_when_the_homepage_stands_in():
    row = _run(on_05(SPECS["ON-05"], _ctx([_page("https://www.acme.com/", page_type="home"), _page("https://www.acme.com/about")])))
    assert [x["url"] for x in row[CHECKED_LINKS]] == ["https://www.acme.com/"]
    assert "stands in" in row[CHECKED_LINKS][0]["why"]


def test_on19_leaves_out_utility_pages():
    pages = [_page("https://www.acme.com/", page_type="home"), _page("https://www.acme.com/careers", page_type="utility"),
             _page("https://www.acme.com/services/data", page_type="service", text="Different words entirely for this one.")]
    row = _run(on_19(SPECS["ON-19"], _ctx(pages)))
    assert {x["url"] for x in row[CHECKED_LINKS]} == {"https://www.acme.com/", "https://www.acme.com/services/data"}


def test_off07_lists_the_search_every_result_and_the_sources_opened(monkeypatch):
    items = [{"url": "https://www.reddit.com/r/it/1", "title": "Acme Corp", "snippet": "Acme Corp", "date": None},
             {"url": "https://www.acme.com/blog", "title": "Acme Corp", "snippet": "own site", "date": None},
             {"url": "https://news.example/story", "title": "Acme Corp", "snippet": "not a forum", "date": None}]

    async def search(query, num=10):
        return gs.SearchResult(query, ok=True, items=items, status=200)

    async def page(url, ctx=None):
        return SimpleNamespace(status_code=200, ok=True, text="<title>Acme Corp</title><p>Acme Corp</p>", error=None)

    monkeypatch.setattr(offpage_tools, "google_search", search)
    monkeypatch.setattr(offpage_tools, "crawl_page", page)
    row = _run(offpage.off_07({"parameter_id": "OFF-07", "section": "off_page", "name": "OFF-07", "weight": 1.0},
                              _ctx([_page("https://www.acme.com/")])))
    by_url = {x["url"]: x["why"] for x in row[CHECKED_LINKS]}
    search_page = next(u for u in by_url if u.startswith("https://www.google.com/search"))
    assert by_url[search_page].startswith("Web search run for this check")
    assert by_url["https://www.reddit.com/r/it/1"].startswith("Forum or community thread naming the company: counted")
    assert "Source page opened to confirm it is about the company (result: verified)" in by_url["https://www.reddit.com/r/it/1"]
    assert "did not count" in by_url["https://www.acme.com/blog"]
    assert "did not count" in by_url["https://news.example/story"]


def test_links_ride_on_the_row_and_never_reach_the_evidence_the_model_grades():
    robots = SimpleNamespace(url="https://www.acme.com/robots.txt", result=SimpleNamespace(ok=False), raw="")
    row = _run(run_one(SPECS["TECH-01"], _ctx([_page("https://www.acme.com/")], robots=robots)))
    assert row[CHECKED_LINKS]
    assert CHECKED_LINKS not in row["evidence"]
    assert CHECKED_LINKS not in (row["evidence"].get("curated_input") or "")


def _rows():
    return [
        {"parameter_id": "ON-05", "section": "on_page", "name": "FAQ on service pages", "evidence": {},
         CHECKED_LINKS: [{"url": "https://www.acme.com/services/a", "why": "Service page"}]},
        {"parameter_id": "TECH-01", "section": "technical", "name": "Robots", "evidence": {},
         CHECKED_LINKS: [{"url": "https://www.acme.com/robots.txt", "why": "robots"}]},
        {"parameter_id": "TECH-21", "section": "technical", "name": "hreflang", "status": "N/A",
         "evidence": {"does_not_apply": "the site serves one market"}},
    ]


def _site():
    return _ctx([_page("https://www.acme.com/", page_type="home")])


def test_no_workbook_for_a_site_that_could_not_be_read(tmp_path, monkeypatch):
    """tcs.com refused every request with 403: no workbook, and no folder made just to hold one."""
    monkeypatch.setattr(parameter_links, "SITE_FOLDERS_ROOT", tmp_path / "sites")
    monkeypatch.setattr(parameter_links, "OUTPUT_DIR", tmp_path / "data")
    refused = _page("https://www.acme.com/")
    refused.result.ok = False
    rows = _rows()
    assert parameter_links.save("scan-1", _ctx([refused]), rows, parameter_links.take(rows)) is None
    assert not (tmp_path / "sites").exists()


def test_take_strips_the_links_off_the_rows():
    rows = _rows()
    links = parameter_links.take(rows)
    assert not any(CHECKED_LINKS in r for r in rows)
    assert links["TECH-01"][0]["url"].endswith("robots.txt") and links["TECH-21"] == []


def test_the_workbook_has_the_three_columns_and_every_parameter(tmp_path, monkeypatch):
    monkeypatch.setattr(parameter_links, "SITE_FOLDERS_ROOT", tmp_path / "sites")
    monkeypatch.setattr(parameter_links, "OUTPUT_DIR", tmp_path / "data")
    rows = _rows()
    links = parameter_links.take(rows)
    path = parameter_links.save("scan-1", _site(), rows, links)
    assert path == tmp_path / "sites" / "www.acme.com" / "acme.xlsx"

    wb = load_workbook(path, read_only=True)
    assert wb.sheetnames == ["Checked Links", "Parameters"]
    sheet = list(wb["Checked Links"].iter_rows(values_only=True))
    assert sheet[0] == ("Parameter Name", "Checked Links", "Reason")
    # Technical before on-page, each in number order; a parameter with no links still has a row saying why.
    assert sheet[1:] == [
        ("TECH-01 - Robots", "https://www.acme.com/robots.txt", "robots"),
        ("TECH-21 - hreflang", "(no links checked)", "No links: this check does not apply to this site (the site serves one market)."),
        ("ON-05 - FAQ on service pages", "https://www.acme.com/services/a", "Service page"),
    ]
    summary = list(wb["Parameters"].iter_rows(values_only=True))
    assert summary[0] == ("Parameter Name", "Links Checked", "Reason")
    assert [(r[0], r[1]) for r in summary[1:]] == [("TECH-01 - Robots", 1), ("TECH-21 - hreflang", 0), ("ON-05 - FAQ on service pages", 1)]
    assert "WHAT THIS CHECK READS" in summary[1][2]


def test_a_recheck_keeps_the_links_of_the_parameters_it_did_not_rerun(tmp_path, monkeypatch):
    monkeypatch.setattr(parameter_links, "SITE_FOLDERS_ROOT", tmp_path / "sites")
    monkeypatch.setattr(parameter_links, "OUTPUT_DIR", tmp_path / "data")
    rows = _rows()
    parameter_links.save("scan-1", _site(), rows, parameter_links.take(rows))

    rerun = {"ON-05": [{"url": "https://www.acme.com/services/b", "why": "Service page (re-checked)"}]}
    path = parameter_links.save("scan-1", _site(), _rows(), rerun, merge=True)
    got = [(r[0][:7], r[1]) for r in load_workbook(path, read_only=True)["Checked Links"].iter_rows(min_row=2, values_only=True)]
    assert ("TECH-01", "https://www.acme.com/robots.txt") in got
    assert ("ON-05 -", "https://www.acme.com/services/b") in got
    assert ("ON-05 -", "https://www.acme.com/services/a") not in got

    # The sidecar belongs to one audit: a re-check of another audit does not borrow its links.
    saved = json.loads((tmp_path / "data" / "acme-com-parameter-links.json").read_text(encoding="utf-8"))
    assert saved["scan_id"] == "scan-1"
    path = parameter_links.save("scan-2", _site(), _rows(), rerun, merge=True)
    got = [r[1] for r in load_workbook(path, read_only=True)["Checked Links"].iter_rows(min_row=2, values_only=True)]
    assert "https://www.acme.com/robots.txt" not in got


def test_a_locked_workbook_never_fails_the_audit(tmp_path, monkeypatch):
    monkeypatch.setattr(parameter_links, "SITE_FOLDERS_ROOT", tmp_path / "sites")
    monkeypatch.setattr(parameter_links, "OUTPUT_DIR", tmp_path / "data")

    def locked(*_a, **_k):
        raise PermissionError("open in Excel")

    monkeypatch.setattr(parameter_links, "_write", locked)
    rows = _rows()
    assert parameter_links.save("scan-1", _site(), rows, parameter_links.take(rows)) is None
